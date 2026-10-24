from pathlib import Path
import sys
import tempfile
import unittest

BASE = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE))
from common import inventory
from docx_assemble_core import assemble


class AssembleTests(unittest.TestCase):
    def setUp(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx optional dependency is not installed")
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.docx"
        doc = Document()
        doc.add_heading("Later", level=1)
        doc.add_paragraph("Second article body")
        doc.add_paragraph("02.02.2024")
        doc.add_heading("Earlier", level=1)
        doc.add_paragraph("First article body")
        doc.add_paragraph("01.02.2024")
        doc.save(self.source)
        self.parts = [{"title": "Origins", "articles": [2]},
                      {"title": "Development", "articles": [1]}]

    def test_dry_run_and_exclusive_output(self):
        output = Path(self.temp.name) / "compiled.docx"
        source_bytes = self.source.read_bytes()
        preview = assemble(str(self.source), self.parts, output_path=str(output), contents_title="İçindekiler")
        self.assertTrue(preview["dry_run"])
        self.assertFalse(output.exists())
        with self.assertRaisesRegex(ValueError, "expected_source_sha256"):
            assemble(str(self.source), self.parts, output_path=str(output),
                     expected_source_sha256="stale", dry_run=False)
        result = assemble(str(self.source), self.parts, output_path=str(output),
                          expected_source_sha256=preview["source_sha256"], dry_run=False,
                          contents_title="İçindekiler")
        self.assertEqual(result["article_count"], 2)
        self.assertEqual(self.source.read_bytes(), source_bytes)
        before = inventory(str(self.source))["items"]
        after = inventory(str(output), article_level=2)["items"]
        from docx import Document
        self.assertEqual(Document(output).paragraphs[0].text, "İçindekiler")
        self.assertEqual([item["title"] for item in after], ["Earlier", "Later"])
        self.assertEqual([item["body_sha256"] for item in after],
                         [before[1]["body_sha256"], before[0]["body_sha256"]])
        with self.assertRaisesRegex(ValueError, "already exists"):
            assemble(str(self.source), self.parts, output_path=str(output),
                     expected_source_sha256=preview["source_sha256"], dry_run=False)

    def test_invalid_plan_cannot_write(self):
        output = Path(self.temp.name) / "invalid.docx"
        with self.assertRaisesRegex(ValueError, "Invalid chapter plan"):
            assemble(str(self.source), [{"title": "Partial", "articles": [1]}],
                     output_path=str(output), dry_run=False)
        self.assertFalse(output.exists())

    def test_existing_contents_requires_review(self):
        from docx import Document
        doc = Document(self.source)
        doc.paragraphs[0].insert_paragraph_before("İçindekiler")
        doc.save(self.source)
        preview = assemble(str(self.source), self.parts)
        with self.assertRaisesRegex(ValueError, "Existing contents list"):
            assemble(str(self.source), self.parts,
                     output_path=str(Path(self.temp.name) / "compiled.docx"),
                     expected_source_sha256=preview["source_sha256"], dry_run=False)

    def test_automatic_toc_is_replaced(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document(self.source)
        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        para = OxmlElement("w:p")
        run = OxmlElement("w:r")
        field = OxmlElement("w:instrText")
        field.text = ' TOC \\o "1-3" '
        run.append(field)
        para.append(run)
        content.append(para)
        sdt.append(content)
        doc.element.body.insert(0, sdt)
        doc.save(self.source)
        output = Path(self.temp.name) / "compiled.docx"
        preview = assemble(str(self.source), self.parts, output_path=str(output))
        result = assemble(str(self.source), self.parts, output_path=str(output),
                          expected_source_sha256=preview["source_sha256"], dry_run=False)
        self.assertIn("Replaced source automatic TOC", result["warnings"])
        assembled = Document(output)
        self.assertFalse(any(element.tag == qn("w:sdt") for element in assembled.element.body))


if __name__ == "__main__":
    unittest.main()
